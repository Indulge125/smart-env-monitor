from pathlib import Path
from shutil import copy2
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\develop\programfirst")
TEMPLATE = Path(r"D:\BaiduNetdiskDownload") / "项目计划书.docx"
OUT_DIR = ROOT / "参赛材料"
ASSET_DIR = OUT_DIR / "assets_new_template"
OUTPUT = OUT_DIR / "智境-项目计划书-新模板内容修改版.docx"


def font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def text_size(draw, text, fnt):
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw, text, fnt, max_width):
    lines = []
    for raw in text.split("\n"):
        current = ""
        for ch in raw:
            candidate = current + ch
            if text_size(draw, candidate, fnt)[0] <= max_width or not current:
                current = candidate
            else:
                lines.append(current)
                current = ch
        if current:
            lines.append(current)
    return lines or [""]


def draw_centered(draw, box, text, fnt, fill=(20, 32, 48), line_gap=8):
    x1, y1, x2, y2 = box
    lines = wrap_text(draw, text, fnt, x2 - x1 - 32)
    line_heights = [text_size(draw, line, fnt)[1] for line in lines]
    total_h = sum(line_heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, h in zip(lines, line_heights):
        w, _ = text_size(draw, line, fnt)
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=fnt, fill=fill)
        y += h + line_gap


def rounded_box(draw, box, fill, outline, title, body=None, title_size=34, body_size=26):
    draw.rounded_rectangle(box, radius=24, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = box
    if body:
        draw_centered(draw, (x1 + 12, y1 + 20, x2 - 12, y1 + 78), title, font(title_size, True))
        draw_centered(draw, (x1 + 12, y1 + 86, x2 - 12, y2 - 18), body, font(body_size))
    else:
        draw_centered(draw, box, title, font(title_size, True))


def arrow(draw, start, end, color=(34, 107, 221), width=8):
    draw.line((start, end), fill=color, width=width)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex >= sx else -1
        points = [(ex, ey), (ex - 24 * direction, ey - 14), (ex - 24 * direction, ey + 14)]
    else:
        direction = 1 if ey >= sy else -1
        points = [(ex, ey), (ex - 14, ey - 24 * direction), (ex + 14, ey - 24 * direction)]
    draw.polygon(points, fill=color)


def make_system_architecture(path):
    img = Image.new("RGB", (1600, 900), "#f6f9fc")
    d = ImageDraw.Draw(img)
    d.text((70, 60), "智境系统总体架构", font=font(50, True), fill="#0f2948")
    d.text((72, 124), "感知层 -> 控制层 -> 执行层/通信层 -> 展示层", font=font(28), fill="#5b6f8a")

    boxes = [
        ((80, 255, 360, 430), "#e8f5ff", "#0284c7", "感知层", "光敏 / 热敏\nADC 原始量"),
        ((470, 230, 760, 460), "#dcfce7", "#16a34a", "控制层", "STM32F103C8T6\nFreeRTOS\n任务调度 + 队列"),
        ((890, 210, 1210, 355), "#fef3c7", "#d97706", "执行层", "OLED / LED / 蜂鸣器\nSG90 舵机"),
        ((890, 470, 1210, 615), "#ede9fe", "#7c3aed", "通信层", "USART1 -> 4G DTU\nJSON 上报"),
        ((1320, 320, 1535, 540), "#f8fafc", "#475569", "展示层", "OLED\n本地演示台\n银尔达平台"),
    ]
    for box, fill, outline, title, body in boxes:
        rounded_box(d, box, fill, outline, title, body)

    arrow(d, (360, 342), (470, 342))
    arrow(d, (760, 342), (890, 282), "#16a34a")
    arrow(d, (760, 342), (890, 540), "#16a34a")
    arrow(d, (1210, 540), (1320, 430), "#7c3aed")

    d.rounded_rectangle((250, 710, 1380, 800), radius=18, fill="#ffffff", outline="#cbd5e1", width=2)
    draw_centered(
        d,
        (270, 715, 1360, 795),
        "数据链路：PA0/PA1 采样 -> 16 点滑动均值 -> 温度/亮度换算 -> OLED/舵机/串口 JSON -> 展示页面与平台",
        font(28),
        fill="#0f172a",
    )
    img.save(path, quality=95)


def make_task_architecture(path):
    img = Image.new("RGB", (1600, 900), "#ffffff")
    d = ImageDraw.Draw(img)
    d.text((70, 62), "FreeRTOS 任务分层与协作", font=font(48, True), fill="#0f2948")
    d.text((72, 124), "CMSIS-RTOS v2，6 个线程，2 个消息队列，10KB Heap", font=font(26), fill="#64748b")

    specs = [
        ((80, 230, 440, 370), "#e0f2fe", "#0284c7", "SensorTask", "1s 采样 ADC\n5s 输出环境 JSON"),
        ((580, 205, 940, 350), "#dcfce7", "#16a34a", "DisplayTask", "OLED 显示\n亮度/温度/模式"),
        ((580, 430, 940, 575), "#fef3c7", "#d97706", "ControlTask", "阈值判断\n舵机/LED/报警"),
        ((80, 495, 440, 635), "#f8fafc", "#475569", "KeyTask", "按键去抖\n模式/阈值设置"),
        ((1085, 265, 1455, 430), "#ede9fe", "#7c3aed", "WifiTask / DTU", "USART1 JSON\n上云与命令解析"),
        ((1085, 585, 1455, 720), "#f8fafc", "#94a3b8", "defaultTask", "系统空闲维护"),
    ]
    for item in specs:
        rounded_box(d, item[0], item[1], item[2], item[3], item[4])

    arrow(d, (440, 300), (580, 278))
    arrow(d, (440, 300), (580, 500))
    arrow(d, (440, 565), (580, 505), "#475569")
    arrow(d, (940, 505), (1085, 348), "#d97706")

    d.rounded_rectangle((260, 765, 1260, 835), radius=18, fill="#f8fafc", outline="#cbd5e1", width=2)
    draw_centered(d, (270, 768, 1250, 832), "解耦价值：采集、显示、控制、通信互不阻塞，适合竞赛演示与后续工程扩展。", font(26))
    img.save(path, quality=95)


def make_data_flow(path):
    img = Image.new("RGB", (1600, 900), "#f8fafc")
    d = ImageDraw.Draw(img)
    d.text((70, 62), "环境数据采集与展示流程", font=font(48, True), fill="#0f2948")
    d.text((72, 124), "从传感器采样到本地显示、执行调控与远程管理的闭环", font=font(26), fill="#64748b")

    steps = [
        ("传感器 AO\n光照/温度\n模拟电压", "#ffffff"),
        ("ADC1\nPA0 / PA1\n12 位采样", "#ffffff"),
        ("滤波换算\n16 点均值\n亮度% / 温度°C", "#ffffff"),
        ("实时控制\n阈值判断\n舵机/LED/蜂鸣器", "#ffffff"),
        ("JSON 输出\n[SENSOR]\nUSART2/DTU", "#ffffff"),
        ("展示管理\nWeb 仪表盘\nOLED / 平台", "#ffffff"),
    ]
    x = 70
    box_w = 220
    gap = 58
    boxes = []
    for title, fill in steps:
        box = (x, 300, x + box_w, 455)
        boxes.append(box)
        rounded_box(d, box, fill, "#cbd5e1", title, title_size=27)
        x += box_w + gap
    for a, b in zip(boxes, boxes[1:]):
        arrow(d, (a[2], 377), (b[0], 377), "#2563eb", 6)

    d.rounded_rectangle((170, 640, 1430, 745), radius=18, fill="#0f172a", outline="#0f172a")
    draw_centered(
        d,
        (190, 650, 1410, 735),
        '{"light":68,"temp":27,"mode":0,"servo":2500,"sw1":1,"rssi":18}',
        font(31),
        fill="#f8fafc",
    )
    d.text((210, 790), "展示建议：遮挡光敏电阻、轻微加热热敏模块，观察 OLED、舵机、页面数值和趋势同步变化。", font=font(25), fill="#334155")
    img.save(path, quality=95)


def ensure_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    paths = {
        "system": ASSET_DIR / "zhijing_system_architecture.png",
        "task": ASSET_DIR / "zhijing_freertos_task.png",
        "flow": ASSET_DIR / "zhijing_data_flow.png",
    }
    make_system_architecture(paths["system"])
    make_task_architecture(paths["task"])
    make_data_flow(paths["flow"])
    return paths


def replace_paragraph_text(paragraph, text):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    if text:
        paragraph.add_run(text)


def replace_drawing(paragraph, image_path, width=5.75):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_cell(cell, text):
    cell.text = text


def set_table(table, rows):
    for r_idx, row_values in enumerate(rows):
        if r_idx >= len(table.rows):
            break
        for c_idx, value in enumerate(row_values):
            if c_idx >= len(table.rows[r_idx].cells):
                break
            set_cell(table.rows[r_idx].cells[c_idx], value)


def update_paragraphs(doc):
    replacements = {
        41: "随着室内学习、办公与实验空间对舒适度、节能性和安全性的要求不断提高，传统依靠人工观察或单一传感器显示的环境管理方式已难以满足连续监测、及时调控和远程管理的需求。本项目面向教室、宿舍、实验室、小型机房和创客空间等典型室内场景，设计并实现“智境 — 基于 FreeRTOS 的室内环境智能感知与调控系统”。系统以 STM32F103C8T6 为主控，围绕光照、温度等环境变量构建感知、处理、控制、显示与通信一体化的物联网终端。",
        42: "项目在嵌入式端采用 FreeRTOS 进行多任务调度，将传感器采集、OLED 显示、阈值判断、舵机/声光执行、按键交互和串口通信拆分为独立任务，并通过消息队列完成数据传递。传感器数据经 ADC 采样、16 点滑动均值滤波和标定换算后，生成温度、亮度、模式、执行状态等结构化数据；本地通过 OLED 和 Web Serial 演示台实时展示，远程通过 4G DTU 与银尔达 IoT/DTU 管理平台形成设备在线管理与数据通道。",
        43: "与普通单片机环境监测装置相比，本作品突出实时操作系统分层、边缘侧自动调控、4G 远程通信和可视化展示联动；与单纯平台接入演示相比，本作品能够在板端完成真实传感器采样与执行控制，形成“环境变化—数据采集—状态判断—执行调控—页面展示—平台管理”的闭环。项目硬件模块清晰、成本可控、可扩展性强，具备竞赛展示、教学实验和小型室内空间智能化改造的应用价值。",
        44: "关键词：室内环境监测；STM32F103；FreeRTOS；4G DTU；智能调控；物联网",
        45: "",
        50: "With the increasing demand for comfort, energy efficiency and safety in indoor learning, office and laboratory spaces, manual observation and single-sensor display are no longer sufficient for continuous environmental management. This project, named “Zhijing — an Indoor Environment Intelligent Sensing and Regulation System Based on FreeRTOS”, implements an IoT terminal integrating sensing, control, display and communication on the STM32F103C8T6 microcontroller.",
        51: "The embedded firmware adopts FreeRTOS to decouple sensor acquisition, OLED display, threshold judgment, actuator control, key interaction and serial communication into independent tasks. ADC data from light and temperature sensors is filtered by a 16-point moving average and converted into meaningful environmental values. The system then packages temperature, light, mode and actuator status into structured JSON data for local visualization and 4G DTU based remote device management.",
        52: "Compared with traditional microcontroller monitoring demos, the proposed system provides a complete closed loop from real sensor acquisition to edge-side regulation, local human-machine interaction and cloud-oriented device management. The modular hardware and software design makes it suitable for competition demonstration, teaching practice and further expansion to humidity, CO2, human presence and energy-saving control scenarios.",
        53: "Keywords: indoor environment monitoring; STM32F103; FreeRTOS; 4G DTU; intelligent regulation; Internet of Things",
        57: "第一章 设计需求分析",
        58: "1.1  项目背景与应用场景",
        59: "室内环境质量直接影响学习效率、办公体验、设备运行稳定性和能源消耗。教室、宿舍、实验室和小型机房等空间通常存在光照变化快、局部温升明显、人工调节滞后和设备运行状态难以集中管理等问题。若能利用低成本嵌入式终端对环境数据进行连续采集，并在边缘端快速完成判断与调控，就可以显著提升空间管理的自动化水平。",
        60: "本项目以“真实传感器采集 + FreeRTOS 实时调度 + 本地/远程可视化”为核心思路，将光照、温度等室内环境变量转化为可显示、可调控、可上传的结构化数据。系统适合用于物联网课程实验、竞赛演示、智能教室原型、小型温控/光控场景和室内节能控制验证。",
        61: "1.2  现有方案痛点与需求定位",
        62: "常见环境监测方案要么停留在单传感器数码管/OLED 显示层面，要么依赖 WiFi 局域网或上位机持续连接，难以兼顾真实采集、实时调度、执行控制和远程管理。部分平台接入案例侧重通信链路展示，缺少板端环境变化与执行器动作之间的闭环关系。本项目定位为可展示、可扩展、可工程化迁移的室内环境智能终端，重点解决多任务并发、数据稳定、现场交互和 4G 管理通道集成四类需求。",
        65: "1.3  设计目标与技术指标",
        66: "本作品的设计目标不是简单显示一个温度或亮度数值，而是打造一套完整的室内环境感知与调控原型：前端通过光敏、热敏传感器采集真实模拟量；中端由 STM32 与 FreeRTOS 完成滤波、换算、阈值判断和任务协作；执行端通过 SG90 舵机、LED 和蜂鸣器模拟窗帘、通风或报警动作；展示端通过 OLED、本地网页和银尔达平台呈现系统状态。",
        68: "系统设计强调响应及时、数据可读、链路清晰和演示稳定。采集周期、显示刷新、串口上报和远程管理各自独立运行，避免某一环节阻塞影响整体效果；数据字段采用可扩展 JSON 结构，便于后续对接更多传感器和产品物模型。",
        71: "1.4  总体技术路线",
        72: "总体技术路线可概括为“传感器输入 — ADC 采样 — 滑动滤波 — 环境量换算 — 阈值判断 — 执行控制 — 数据封装 — 本地/远程展示”的闭环架构。该路线以板端真实环境数据为基础，通过实时操作系统保证多模块协同运行。",
        73: "在硬件层面，PA0/PA1 分别采集光照和温度模拟量，PA8 输出舵机 PWM，PA6/PA7/PA11 用于蜂鸣器和指示灯，PB8/PB9 驱动 OLED，PA9/PA10 与 4G DTU 通信，PA2/PA3 用于本地调试与网页展示。软件层面采用 HAL + CMSIS-RTOS v2 组织任务，形成清晰的数据流和控制流。",
        76: "图 1-1  智境系统总体架构示意图",
        77: "本章小结：本项目需求来自室内环境连续监测与自动调控场景，重点解决真实采集、多任务协同、执行联动和远程管理问题。系统以 STM32F103C8T6 与 FreeRTOS 为核心，形成适合竞赛演示和后续扩展的物联网终端方案。",
        79: "第二章 特色与创新",
        80: "2.1  作品介绍",
        81: "智境是一套面向室内空间的智能环境感知与调控系统。作品通过光敏、热敏传感器获取环境状态，利用 STM32F103C8T6 对数据进行滤波和换算，并依据工作模式与阈值策略驱动舵机、LED、蜂鸣器等执行器，从而模拟自动遮光、通风调节、环境提醒等应用场景。",
        82: "作品展示时，可通过遮挡光敏模块、轻微加热热敏模块和按键切换模式来触发系统状态变化，观众能够同时看到 OLED 数值、执行器动作、本地仪表盘趋势和平台设备在线信息，形成直观的物联网闭环展示效果。",
        83: "2.2  关键技术创新点",
        86: "2.3  项目应用价值",
        87: "在教学与竞赛价值方面，本项目覆盖 STM32 外设、ADC 采样、PWM 控制、OLED 显示、FreeRTOS 任务调度、串口通信、JSON 数据封装和平台接入等典型物联网知识点，能够完整体现嵌入式软件、硬件电路和云端管理的协同设计能力。",
        89: "在应用价值方面，系统可迁移到教室节能管理、宿舍舒适度提醒、实验室设备环境监测、小型机房温升预警和创客空间自动调光等场景。通过替换传感器模块，还可扩展到湿度、CO2、烟雾、人体存在和空气质量等指标。",
        90: "在推广价值方面，作品采用模块化接口和标准化数据字段，便于后续在嘉立创 EDA 中设计一块集成底板，将传感器、DTU、OLED、按键和执行器统一插接，形成更接近产品形态的竞赛样机。",
        91: "本章小结：作品特色不只在于完成环境数据采集，而在于将实时操作系统、边缘控制、4G 通信、可视化展示和硬件集成思路串成完整闭环。",
        93: "第三章 功能设计",
        94: "3.1  系统总体功能框架",
        95: "系统功能围绕室内环境监测与调控展开，整体包括传感器采集模块、数据滤波与换算模块、FreeRTOS 任务调度模块、阈值判断与执行控制模块、OLED 人机交互模块、串口/4G 通信模块、本地仪表盘模块和平台设备管理模块。各模块通过结构化数据连接，便于调试和扩展。",
        97: "图 3-1  环境数据采集与展示流程示意图",
        98: "3.2  传感器采集与数据预处理",
        99: "传感器采集模块负责将光照和温度变化转换为 ADC 原始值。为了减少模拟传感器抖动和瞬时噪声，系统对采样值进行滑动均值滤波，并按设定公式换算为亮度百分比和温度估计值。预处理后的数据既用于 OLED 显示，也用于阈值判断和串口 JSON 输出。",
        100: "表 3-1  环境数据预处理环节",
        102: "3.3  FreeRTOS 多任务调度模块",
        103: "系统采用 FreeRTOS 将核心功能拆分为多个任务。SensorTask 负责周期采样与数据封装，DisplayTask 负责 OLED 刷新，ControlTask 负责模式和阈值判断，KeyTask 负责按键去抖与参数设置，通信任务负责串口数据输出和远程链路维护，defaultTask 负责系统空闲维护。",
        105: "多任务架构的优势在于采集、显示、控制和通信互不阻塞。例如，即使通信链路正在进行数据发送，传感器采样和本地执行判断仍能按周期运行，保证现场演示效果稳定。",
        106: "任务之间通过消息队列传递环境数据和控制状态，避免多个任务直接争用全局变量。该设计既符合 FreeRTOS 工程实践，也便于后续增加湿度、CO2 或人体检测任务。",
        107: "表 3-2  核心方案选型对比",
        110: "3.4  环境数据滤波与状态判断模块",
        111: "光照和热敏传感器属于模拟量输入，原始值会受到电源波动、接线长度和环境扰动影响。系统通过 16 点滑动均值滤波降低抖动，再将 ADC 数值换算为便于理解的亮度百分比和温度值。",
        112: "状态判断模块根据工作模式和阈值对环境状态进行分类。当亮度低于设定值时，可触发补光或状态提示；当温度超过设定值时，可驱动舵机模拟通风角度变化，并通过 LED 或蜂鸣器给出提示。",
        113: "该模块的设计重点是让环境变化能够被用户直接观察：遮挡光敏电阻后亮度下降，OLED 和页面同步更新；轻微加热热敏模块后温度上升，舵机和报警状态随之变化，从而形成真实、可复现的演示链路。",
        116: "图 3-2  FreeRTOS 任务协作示意图",
        117: "3.5  执行控制与阈值调节模块",
        118: "执行控制模块通过 PWM、GPIO 和状态变量实现室内调控动作。PA8 输出 PWM 控制 SG90 舵机角度，可模拟通风窗、遮阳片或调节机构；LED 和蜂鸣器用于状态提示和报警提醒；按键用于切换工作模式、调整阈值和现场演示。",
        119: "阈值调节采用按键交互与 OLED 显示配合完成，用户不需要重新烧录程序即可改变控制策略。该设计使作品既能展示自动模式，也能展示人工设定和现场调参能力。",
        120: "表 3-3  传感器与执行器接口参数",
        122: "3.6  4G 通信与平台管理模块",
        123: "通信模块通过 USART1 连接 M100M-C2 4G DTU，利用移动网络实现设备在线管理和数据通道扩展。系统上行数据采用 JSON 结构，包含温度、亮度、工作模式、舵机状态、开关量和信号状态等字段，可与银尔达 IoT 产品功能点进行对应，便于后续进行平台规则、告警和可视化页面配置。",
        126: "图 3-3  4G DTU 与平台管理链路示意图",
        127: "3.7  可视化展示与人机交互",
        128: "可视化展示由 OLED、本地网页仪表盘和银尔达平台共同组成。OLED 负责现场快速查看，网页仪表盘通过串口读取板端 JSON 数据并绘制趋势曲线，平台用于设备在线状态、SIM 卡信息、产品功能点和远程管理展示。",
        129: "本地仪表盘适合比赛录制视频：打开网页后连接串口，系统按设定周期接收来自板端的真实 JSON 数据，温度、亮度、模式和舵机状态会随传感器变化更新。页面同时显示串口日志和最新上报内容，便于展示数据来源。",
        130: "平台侧可展示设备在线、产品信息、功能点定义、通信参数和设备日志，使作品从单板实验扩展为具备远程管理能力的物联网系统。",
        133: "图 3-4  项目数据闭环与管理图",
        134: "界面设计遵循“实时、清晰、可验证”的原则：顶部显示连接状态和刷新周期，中部以卡片呈现温度、亮度、模式、执行器和信号状态，下方以趋势图展示数据变化。页面不承担业务造数，而是解析串口收到的板端数据，保证演示可信。",
        135: "在交互流程上，用户先连接串口并选择波特率，再通过改变光照或温度触发板端状态变化，最后观察 OLED、执行器和网页趋势同步更新。该流程简单直观，适合答辩现场和讲解视频拍摄。",
        137: "图 3-5  本地实时监测页面结构示意图",
        139: "图 3-6  设备管理与数据展示结构示意图",
        141: "第四章 系统实现",
        142: "4.1  感知层技术实现",
        143: "感知层由光敏传感器和热敏传感器组成，分别接入 STM32 的 PA0/ADC1_IN0 和 PA1/ADC1_IN1。系统周期读取 ADC 原始值，并通过滤波和换算得到亮度百分比与温度估计值。该层直接面向真实环境变化，是后续显示、控制和上报的基础。",
        146: "4.2  实时任务层技术实现",
        147: "实时任务层采用 STM32 HAL 与 CMSIS-RTOS v2 API 进行开发。系统创建传感器、显示、控制、按键、通信和默认维护任务，并配置 10KB FreeRTOS Heap 与消息队列。任务优先级和周期根据功能重要性设置，使采集与控制保持稳定。",
        148: "在程序实现上，ADC 采集、OLED 刷新、PWM 输出和串口收发均由独立模块封装，主逻辑通过任务调度协同运行。该结构降低了代码耦合度，也便于后续新增传感器或调整通信协议。",
        152: "4.3  数据层与传输层实现",
        153: "数据层采用统一 JSON 字段描述环境状态，例如 light、temp、mode、servo、sw1、in1、rssi 等。传输层通过 USART2 输出本地调试数据，通过 USART1 与 4G DTU 进行远程通信。银尔达平台侧可在产品功能点中建立对应标识，使设备状态具备进一步展示、规则判断和扩展管理的基础。",
        154: "表 4-1  物联网数据字段设计",
        156: "4.4  控制层与可视化应用实现",
        157: "控制层根据温度、亮度和工作模式驱动舵机、LED 与蜂鸣器，并将执行结果同步写入状态数据。可视化应用负责解析串口 JSON、刷新数值卡片、绘制实时趋势和记录日志；OLED 则提供脱离电脑后的现场状态显示。三者共同提升了作品展示的完整性和可理解性。",
        161: "4.5  测试结果与性能分析",
        162: "原型测试表明，系统能够稳定读取光敏与热敏传感器数据，并在 OLED 和本地仪表盘中同步呈现；遮挡光敏模块、改变环境温度或按键切换模式时，数值、状态和执行器动作均能形成可观察的联动效果。",
        163: "通信与平台展示方面，4G DTU 可完成设备上线、SIM 卡管理、通信参数配置和设备日志查看，为作品提供了远程管理入口。结合本地串口仪表盘，系统能够同时满足现场演示和平台化拓展两类需求。",
        166: "本章小结：系统实现采用分层架构，感知层负责真实采样，实时任务层负责调度，控制层负责联动执行，数据层负责结构化输出，展示层负责本地和远程呈现。比赛展示重点建议放在“传感器变化 + OLED/执行器响应 + 网页趋势 + 平台在线”四个环节。",
        168: "第五章 其他内容",
        169: "5.1  工业设计与外观方案",
        170: "作品外观可设计为“桌面环境感知终端 + 远程管理平台”两部分。终端采用一块集成底板承载 STM32 最小系统、DTU、传感器接口、OLED、按键、LED、蜂鸣器、舵机接口、CH340 调试口和 SWD 下载口。后续可使用嘉立创 EDA 绘制底板，使模块插接更牢固、走线更清晰、展示更像完整产品。",
        172: "5.2  项目进度计划",
        174: "项目已完成核心硬件连接、FreeRTOS 软件框架、传感器采集、OLED 显示、执行器控制、本地串口展示和 4G DTU 平台接入配置。后续重点是优化外观底板、整理演示流程、完善 PPT 与讲解视频，使参赛材料与实物演示保持一致。",
        176: "5.3  成本预算与可推广性",
        178: "系统采用常见 STM32、基础传感器和低成本执行器，单套样机成本可控；若批量制作底板，可进一步减少杜邦线和连接器成本。通过更换传感器和调整 JSON 字段，系统可快速迁移到更多室内监测与节能控制场景。",
        179: "5.4  前景展望与优化方向",
        180: "未来可从四个方向继续完善：一是加入湿度、CO2、PM2.5、人体存在等多源传感器；二是完善嘉立创 EDA 一体化底板和外壳设计；三是增加平台规则引擎、告警推送和历史数据分析；四是优化低功耗策略和供电保护，使系统更接近可长期部署的产品形态。",
        183: "本章小结：工业设计建议突出“桌面环境感知终端 + 远程管理平台”的完整形态。",
        184: "后续升级重点可放在多源传感、集成底板、平台可视化、低功耗和外壳产品化。",
        186: "参考文献",
        187: "[1] STMicroelectronics. STM32F103x8/B Datasheet[EB/OL].",
        188: "[2] STMicroelectronics. RM0008 STM32F10xxx Reference Manual[EB/OL].",
        189: "[3] Arm. Cortex-M3 Technical Reference Manual[EB/OL].",
        190: "[4] FreeRTOS. FreeRTOS Kernel Developer Documentation[EB/OL].",
        191: "[5] STMicroelectronics. STM32Cube HAL and LL Drivers User Manual[EB/OL].",
        192: "[6] 银尔达. M100M-C2 4G DTU 模块产品资料与平台使用说明[EB/OL].",
        193: "[7] OASIS. MQTT Version 3.1.1 Standard[EB/OL].",
        194: "[8] MDN Web Docs. Web Serial API Documentation[EB/OL].",
        195: "[9] 常用传感器、OLED、SG90 舵机及嵌入式实验模块产品手册[EB/OL].",
    }
    for idx, text in replacements.items():
        if idx < len(doc.paragraphs):
            replace_paragraph_text(doc.paragraphs[idx], text)


def update_tables(doc):
    set_table(
        doc.tables[0],
        [
            ["需求类别", "具体需求", "项目响应方案"],
            ["实时性", "连续采集室内光照、温度并快速刷新状态", "FreeRTOS 多任务调度，采集、显示、控制和通信并行运行"],
            ["真实性", "展示数据应来自板端传感器而非静态演示", "PA0/PA1 ADC 采样，16 点滑动均值滤波后生成环境数据"],
            ["可调控", "环境异常时能够产生执行动作和现场提示", "SG90 舵机、LED、蜂鸣器与阈值策略联动"],
            ["可展示", "参赛演示需要直观呈现数值、趋势和设备状态", "OLED + 本地网页仪表盘 + 银尔达平台多端展示"],
            ["可扩展", "后续可增加更多环境指标和平台功能", "模块化接口、JSON 字段和产品功能点可扩展"],
        ],
    )
    set_table(
        doc.tables[1],
        [
            ["指标项", "目标值/设计要求", "说明"],
            ["主控平台", "STM32F103C8T6，72MHz Cortex-M3", "适合低成本室内物联网终端"],
            ["系统架构", "FreeRTOS / CMSIS-RTOS v2", "6 个任务、2 个消息队列、10KB Heap"],
            ["采集对象", "光照、温度两类环境量", "可扩展湿度、CO2、空气质量等传感器"],
            ["数据处理", "ADC 12 位采样 + 16 点滑动均值", "降低模拟量抖动，提高显示稳定性"],
            ["通信展示", "USART1 连接 4G DTU，USART2 本地调试", "兼顾平台管理和现场演示"],
            ["执行控制", "舵机 PWM、LED、蜂鸣器、按键交互", "形成环境变化到动作响应的闭环"],
        ],
    )
    set_table(
        doc.tables[2],
        [
            ["创新点", "具体内容", "项目价值"],
            ["RTOS 分层协作", "将采集、显示、控制、按键和通信拆分为独立任务，通过队列传递数据。", "相比顺序循环程序，响应更稳定，结构更利于扩展。"],
            ["边缘侧自动调控", "板端根据真实传感器数据完成阈值判断并驱动舵机、LED、蜂鸣器。", "不依赖上位机即可形成现场智能控制能力。"],
            ["多端可视化展示", "OLED、本地网页仪表盘和银尔达平台共同呈现设备状态。", "便于比赛讲解，也便于体现物联网完整链路。"],
            ["标准化数据格式", "使用 light、temp、mode、servo 等字段封装 JSON 数据。", "方便平台功能点映射和后续增加传感器。"],
            ["模块化硬件集成", "各模块接口清晰，可进一步设计嘉立创 EDA 集成底板。", "提升样机可靠性和产品化展示效果。"],
        ],
    )
    set_table(
        doc.tables[3],
        [
            ["预处理环节", "设计内容", "作用"],
            ["ADC 采样", "周期读取 PA0、PA1 模拟量", "获取光照和温度原始数据"],
            ["滑动滤波", "保存最近 16 次采样并求均值", "降低传感器抖动和瞬时噪声"],
            ["量纲换算", "将 ADC 值换算为亮度百分比和温度估计值", "使数据便于展示和判断"],
            ["状态封装", "组合 light、temp、mode、servo 等字段", "为 OLED、网页和平台提供统一数据源"],
            ["异常判断", "与阈值比较后更新报警和执行器状态", "实现环境调控闭环"],
        ],
    )
    set_table(
        doc.tables[4],
        [
            ["方案", "实时性", "复杂度", "扩展性", "适用建议"],
            ["裸机轮询", "中", "低", "一般", "适合简单入门演示，复杂通信时易阻塞"],
            ["FreeRTOS 多任务", "高", "中", "高", "本项目推荐，便于采集、显示、控制和通信并行"],
            ["仅本地显示", "高", "低", "一般", "适合实验验证，但平台展示能力弱"],
            ["本地 + 4G 平台", "高", "中", "高", "本项目采用，兼顾现场演示和远程管理"],
        ],
    )
    set_table(
        doc.tables[5],
        [
            ["模块/接口", "连接引脚", "说明"],
            ["光敏传感器 AO", "PA0 / ADC1_IN0", "采集光照模拟量"],
            ["热敏传感器 AO", "PA1 / ADC1_IN1", "采集温度相关模拟量"],
            ["SG90 舵机", "PA8 / TIM1_CH1", "PWM 控制角度，模拟调控机构"],
        ],
    )
    set_table(
        doc.tables[6],
        [
            ["设备/输入", "功能", "部署建议"],
            ["STM32F103C8T6 核心板", "运行 FreeRTOS 与控制逻辑", "作为系统主控"],
            ["光敏/热敏传感器", "获取室内环境数据", "布置在能够感知光照和热源变化的位置"],
            ["OLED 与按键", "本地显示和人机交互", "放在前面板，便于演示"],
            ["4G DTU 模块", "远程通信和设备在线管理", "注意天线、SIM 卡和供电稳定"],
            ["本地网页仪表盘", "展示趋势、日志和最新 JSON", "适合录制讲解视频"],
        ],
    )
    set_table(
        doc.tables[7],
        [
            ["阶段", "实现内容", "输出结果"],
            ["需求分析", "确定室内环境监测、调控和平台展示目标", "项目方案与功能框架"],
            ["硬件搭建", "连接 STM32、传感器、OLED、舵机、DTU、按键和指示灯", "可运行原型板"],
            ["RTOS 开发", "创建采集、显示、控制、按键、通信等任务", "稳定的软件框架"],
            ["数据联调", "完成 ADC 滤波、JSON 输出、本地仪表盘解析", "真实数据展示链路"],
            ["平台配置", "配置产品功能点、通信参数和设备管理信息", "远程管理与展示能力"],
            ["参赛完善", "整理计划书、PPT、讲解脚本和视频流程", "完整参赛材料"],
        ],
    )
    set_table(
        doc.tables[8],
        [
            ["字段名", "数据类型", "含义"],
            ["light", "int/float", "亮度百分比，来自光敏传感器换算"],
            ["temp", "int/float", "温度估计值，来自热敏传感器换算"],
            ["mode", "int", "工作模式或控制策略编号"],
            ["servo", "int", "舵机 PWM 或角度状态"],
            ["sw1/in1", "int", "按键、输入或开关量状态"],
            ["rssi", "int", "4G 通信信号强度或平台展示信号字段"],
        ],
    )
    set_table(
        doc.tables[9],
        [
            ["模块", "关键技术", "功能说明"],
            ["传感器采集", "STM32 ADC + 滑动均值滤波", "获取稳定环境数据"],
            ["实时调度", "FreeRTOS / CMSIS-RTOS v2", "组织多任务并发运行"],
            ["显示交互", "OLED 软件 I2C + 按键去抖", "实现现场状态显示与参数设置"],
            ["执行控制", "TIM PWM + GPIO", "驱动舵机、LED 和蜂鸣器"],
            ["通信管理", "USART1/USART2 + 4G DTU", "完成本地调试和远程管理"],
            ["可视化展示", "Web Serial + 趋势图", "展示真实串口数据和运行日志"],
        ],
    )
    set_table(
        doc.tables[10],
        [
            ["测试项", "结果/现象", "项目意义"],
            ["传感器采集", "遮挡光敏或改变热敏环境时数值持续变化", "说明数据来自真实环境输入"],
            ["OLED 显示", "温度、亮度、模式和状态能够周期刷新", "满足脱机现场展示需求"],
            ["执行联动", "阈值触发后舵机、LED、蜂鸣器产生响应", "体现自动调控闭环"],
            ["串口展示", "本地页面可解析 JSON 并绘制趋势", "便于录制参赛演示视频"],
            ["平台管理", "设备、SIM 卡和通信参数可在平台侧管理", "体现物联网远程扩展能力"],
        ],
    )
    set_table(
        doc.tables[11],
        [
            ["组成部分", "外观/结构设计", "作用"],
            ["集成底板", "嘉立创 EDA 绘制模块化插接底板", "减少飞线，提升样机完整度"],
            ["传感器区域", "光敏、热敏模块外露并带固定孔位", "便于演示光照和温度变化"],
            ["显示控制面板", "OLED、按键、LED 集中布置", "方便现场观察和操作"],
            ["通信区域", "DTU、SIM 卡、天线与供电保持间距", "提升远程通信稳定性"],
            ["外壳与支架", "透明上盖或开放式展示结构", "便于评委观察硬件组成"],
        ],
    )
    set_table(
        doc.tables[12],
        [
            ["阶段", "时间安排", "主要任务", "成果"],
            ["第一阶段", "已完成", "需求分析、项目定位、总体方案设计", "项目方案与功能框架"],
            ["第二阶段", "已完成", "硬件模块连接、FreeRTOS 任务框架搭建", "可运行嵌入式原型"],
            ["第三阶段", "已完成", "传感器采集、OLED 显示、执行器联动", "环境感知与调控闭环"],
            ["第四阶段", "已完成", "本地仪表盘、4G DTU、平台配置联调", "展示与管理链路"],
            ["第五阶段", "进行中", "计划书、PPT、讲解视频和集成底板完善", "参赛材料与产品化样机"],
        ],
    )
    set_table(
        doc.tables[13],
        [
            ["项目", "数量", "单价估计", "小计", "说明"],
            ["STM32F103C8T6 核心板", "1", "20-40 元", "20-40 元", "主控与 FreeRTOS 运行平台"],
            ["光敏/热敏传感器", "2", "5-15 元", "10-30 元", "环境数据采集"],
            ["OLED、按键、LED、蜂鸣器", "1 套", "20-50 元", "20-50 元", "显示与交互提示"],
            ["SG90 舵机", "1", "10-20 元", "10-20 元", "模拟通风/遮光执行机构"],
            ["M100M-C2 4G DTU 与 SIM 卡", "1 套", "按已有设备计", "可复用", "远程管理与平台展示"],
            ["PCB 底板、电源与连接件", "1 套", "50-150 元", "50-150 元", "提升展示可靠性和集成度"],
        ],
    )
    set_table(
        doc.tables[14],
        [
            ["优化方向", "可能影响", "改进措施"],
            ["传感器标定", "不同模块输出曲线存在差异", "增加实测标定点，优化换算公式"],
            ["供电稳定性", "DTU 和舵机瞬时电流可能较大", "独立供电、增加滤波电容和保护电路"],
            ["布线可靠性", "杜邦线连接在展示中容易松动", "设计集成底板和锁扣式接口"],
            ["平台展示扩展", "单一字段展示不够丰富", "增加规则引擎、历史曲线和告警配置"],
            ["产品化外观", "裸板展示不够完整", "增加亚克力外壳、传感器窗口和标识面板"],
        ],
    )


def update_images(doc, image_paths):
    drawings = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
    sequence = [
        image_paths["system"],
        image_paths["flow"],
        image_paths["task"],
        image_paths["system"],
        image_paths["flow"],
        image_paths["task"],
        image_paths["system"],
    ]
    for paragraph, image_path in zip(drawings, sequence):
        replace_drawing(paragraph, image_path)


def patch_cached_toc(docx_path):
    replacements = {
        "3.2  视频采集与图像预处理": "3.2  传感器采集与数据预处理",
        "3.2 视频采集与图像预处理": "3.2 传感器采集与数据预处理",
        "3.3 YOLOv8 ": "3.3 FreeRTOS ",
        "车辆检测模块": "多任务调度模块",
        "3.4 IOU-Tracker ": "3.4 环境数据滤波与状态判断",
        "轨迹跟踪模块": "模块",
        "3.5  速度测算与超速判定模块": "3.5  执行控制与阈值调节模块",
        "3.5 速度测算与超速判定模块": "3.5 执行控制与阈值调节模块",
        "速度测算与超速判定模块": "执行控制与阈值调节模块",
        "3.6  超速截图、语音报警与数据闭环": "3.6  4G 通信与平台管理模块",
        "3.6 超速截图、语音报警与数据闭环": "3.6 4G 通信与平台管理模块",
        "超速截图、语音报警与数据闭环": "4G 通信与平台管理模块",
        "3.7  可视化应用与历史查询": "3.7  可视化展示与人机交互",
        "3.7 可视化应用与历史查询": "3.7 可视化展示与人机交互",
        "4.2  算法层技术实现": "4.2  实时任务层技术实现",
        "4.2 算法层技术实现": "4.2 实时任务层技术实现",
        "5.4  前景展望与风险改进": "5.4  前景展望与优化方向",
        "5.4 前景展望与风险改进": "5.4 前景展望与优化方向",
    }

    with ZipFile(docx_path, "r") as zin, NamedTemporaryFile(delete=False, suffix=".docx", dir=docx_path.parent) as tmp:
        tmp_path = Path(tmp.name)
        with ZipFile(tmp, "w", ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith(".xml"):
                    text = data.decode("utf-8", errors="ignore")
                    for old, new in replacements.items():
                        text = text.replace(old, new)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
    tmp_path.replace(docx_path)


def main():
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    copy2(TEMPLATE, OUTPUT)

    image_paths = ensure_assets()
    doc = Document(OUTPUT)
    update_paragraphs(doc)
    update_tables(doc)
    update_images(doc, image_paths)
    doc.save(OUTPUT)
    patch_cached_toc(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
