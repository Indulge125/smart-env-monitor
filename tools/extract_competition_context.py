from __future__ import annotations

from pathlib import Path
from docx import Document


ROOT = Path(r"D:\develop\programfirst")
TEMPLATE = Path(r"D:\qqdownload\智境 — 基于 FreeRTOS 的室内环境智能感知与调控系统项目计划书..docx")
OUT = ROOT / "参赛材料"
OUT.mkdir(exist_ok=True)


def extract_docx_template() -> str:
    doc = Document(str(TEMPLATE))
    lines: list[str] = []
    lines.append(f"TEMPLATE: {TEMPLATE}")
    lines.append(f"PARAGRAPHS: {len(doc.paragraphs)}")
    lines.append(f"TABLES: {len(doc.tables)}")
    lines.append("")

    for idx, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        if not text:
            continue
        lines.append(f"P{idx:03d} [{p.style.name}]: {text}")

    lines.append("")
    lines.append("TABLE_SUMMARY")
    for table_idx, table in enumerate(doc.tables, 1):
        lines.append(f"TABLE {table_idx}: rows={len(table.rows)}, cols={len(table.columns)}")
        for row_idx, row in enumerate(table.rows[:8], 1):
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            lines.append(f"  R{row_idx}: {' | '.join(cells)}")
        if len(table.rows) > 8:
            lines.append("  ...")
    return "\n".join(lines)


def read_text(path: Path, limit: int | None = None) -> str:
    text = path.read_text(encoding="utf-8", errors="ignore")
    if limit is not None and len(text) > limit:
        return text[:limit] + "\n...[TRUNCATED]..."
    return text


def collect_project_context() -> str:
    files = [
        ROOT / "smart_env_monitor" / "smart_env_monitor.ioc",
        ROOT / "smart_env_monitor" / "Core" / "Src" / "freertos.c",
        ROOT / "smart_env_monitor" / "Core" / "Src" / "main.c",
        ROOT / "smart_env_monitor" / "Core" / "Src" / "adc.c",
        ROOT / "smart_env_monitor" / "Core" / "Src" / "gpio.c",
        ROOT / "smart_env_monitor" / "Core" / "Src" / "usart.c",
        ROOT / "smart_env_monitor" / "Core" / "Src" / "tim.c",
        ROOT / "smart_env_monitor" / "Core" / "Inc" / "main.h",
        ROOT / "demo_dashboard" / "index.html",
        ROOT / "demo_dashboard" / "app.js",
        ROOT / "现有模块列表" / "模块列表.md",
        ROOT / "现有模块列表" / "银尔达M00M-C2" / "银尔达模块说明.md",
    ]
    lines: list[str] = []
    for file in files:
        lines.append(f"\n===== {file} =====")
        if file.exists():
            lines.append(read_text(file, 40000))
        else:
            lines.append("MISSING")
    return "\n".join(lines)


def main() -> None:
    (OUT / "模板结构提取.txt").write_text(extract_docx_template(), encoding="utf-8")
    (OUT / "项目上下文提取.txt").write_text(collect_project_context(), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
