from __future__ import annotations

import json
import math
import os
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "参赛材料"
ASSET_DIR = OUT_DIR / "assets"
PRESENTATION_WORKSPACE = ROOT / "outputs" / "manual-zhijing-competition" / "presentations" / "zhijing-iot-contest"
SLIDES_DIR = PRESENTATION_WORKSPACE / "slides"
PREVIEW_DIR = PRESENTATION_WORKSPACE / "preview"
LAYOUT_DIR = PRESENTATION_WORKSPACE / "layout"
QA_DIR = PRESENTATION_WORKSPACE / "qa"

PROJECT_TITLE = "智境 — 基于 FreeRTOS 的室内环境智能感知与调控系统"
PROJECT_SUBTITLE = "物联网设计大赛项目计划书"
TODAY = "2026年6月"


def font_path() -> str:
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return ""


FONT_PATH = font_path()


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    if FONT_PATH:
        return ImageFont.truetype(FONT_PATH, size=size, index=0)
    return ImageFont.load_default()


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    LAYOUT_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)


def set_east_asia_font(run, name: str = "微软雅黑") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_style_font(style, name: str = "Calibri", east_asia: str = "微软雅黑", size: int | None = None, color: str | None = None) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        tag = f"w:{edge}"
        element = tc_mar.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360, indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))


def style_table(table, widths: list[float] | None = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            if widths and col_idx < len(widths):
                cell.width = Inches(widths[col_idx])


def add_run(paragraph, text: str, bold: bool = False, color: str | None = None, size: int | None = None):
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.25 + level * 0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    add_run(paragraph, text)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    add_run(paragraph, text)


def add_note(doc: Document, title: str, body: str, fill: str = "F4F6F9") -> None:
    table = doc.add_table(rows=1, cols=1)
    style_table(table, [6.3])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    add_run(paragraph, title + "：", bold=True, color="1F4D78")
    add_run(paragraph, body)
    doc.add_paragraph()


def setup_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_style_font(styles["Normal"], size=11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Title", 22, "0B2545"),
        ("Subtitle", 12, "4B5563"),
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        set_style_font(styles[name], size=size, color=color)
    styles["Heading 1"].paragraph_format.space_before = Pt(16)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)
    set_style_font(styles["List Bullet"], size=11)
    set_style_font(styles["List Number"], size=11)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(header, PROJECT_TITLE, color="6B7280", size=9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "物联网设计大赛参赛材料 | 生成日期：" + TODAY, color="6B7280", size=9)


def draw_round_box(draw: ImageDraw.ImageDraw, xy, fill, outline, text, font, text_fill="#0F172A", radius=24, align="center"):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=2)
    x1, y1, x2, y2 = xy
    lines = str(text).split("\n")
    line_heights = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        line_heights.append(bbox[3] - bbox[1])
    total_h = sum(line_heights) + (len(lines) - 1) * 8
    y = y1 + (y2 - y1 - total_h) / 2
    for i, line in enumerate(lines):
        bbox = draw.textbbox((0, 0), line, font=font)
        w = bbox[2] - bbox[0]
        if align == "left":
            x = x1 + 26
        else:
            x = x1 + (x2 - x1 - w) / 2
        draw.text((x, y), line, font=font, fill=text_fill)
        y += line_heights[i] + 8


def arrow(draw: ImageDraw.ImageDraw, start, end, color="#2563EB", width=6):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    points = [
        end,
        (end[0] - size * math.cos(angle - math.pi / 6), end[1] - size * math.sin(angle - math.pi / 6)),
        (end[0] - size * math.cos(angle + math.pi / 6), end[1] - size * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(points, fill=color)


def make_architecture_diagram() -> Path:
    path = ASSET_DIR / "system_architecture.png"
    image = Image.new("RGB", (1600, 900), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    title_font = pil_font(46)
    body_font = pil_font(30)
    small_font = pil_font(24)
    draw.text((70, 50), "智境系统总体架构", font=title_font, fill="#0B2545")
    draw.text((72, 112), "感知层 → 控制层 → 执行层/通信层 → 展示层", font=small_font, fill="#64748B")

    boxes = [
        ((80, 230, 360, 420), "#E0F2FE", "#0284C7", "感知层\n光敏 / 热敏\nADC 原始量"),
        ((460, 210, 760, 445), "#DCFCE7", "#16A34A", "STM32F103C8T6\nFreeRTOS\n任务调度 + 队列"),
        ((860, 160, 1180, 320), "#FEF3C7", "#D97706", "执行层\nOLED / LED / 蜂鸣器\nSG90 舵机"),
        ((860, 390, 1180, 550), "#EDE9FE", "#7C3AED", "通信层\nUSART1 → 4G DTU\nJSON 上报"),
        ((1280, 240, 1520, 470), "#F1F5F9", "#475569", "展示层\n本地仪表盘\n银尔达平台"),
    ]
    for xy, fill, outline, text in boxes:
        draw_round_box(draw, xy, fill, outline, text, body_font)
    arrow(draw, (360, 325), (460, 325), "#0284C7")
    arrow(draw, (760, 300), (860, 245), "#16A34A")
    arrow(draw, (760, 360), (860, 470), "#16A34A")
    arrow(draw, (1180, 470), (1280, 355), "#7C3AED")

    draw_round_box(draw, (240, 650, 1360, 785), "#FFFFFF", "#CBD5E1",
                   "真实数据链路：PA0/PA1 采样 → 16 点滑动均值 → 温度/亮度换算 → OLED/舵机/串口 JSON → 展示页面",
                   small_font, align="center", radius=18)
    image.save(path)
    return path


def make_task_diagram() -> Path:
    path = ASSET_DIR / "freertos_task_architecture.png"
    image = Image.new("RGB", (1600, 900), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = pil_font(44)
    body_font = pil_font(27)
    small_font = pil_font(22)
    draw.text((70, 55), "FreeRTOS 任务分层与协作", font=title_font, fill="#0B2545")
    draw.text((70, 112), "CMSIS-RTOS v2，6 个线程，2 个消息队列，10KB Heap", font=small_font, fill="#64748B")

    items = [
        ((80, 210, 430, 350), "#E0F2FE", "#0284C7", "SensorTask\n1s 采样 ADC\n5s 输出真实 JSON"),
        ((80, 430, 430, 570), "#F1F5F9", "#475569", "KeyTask\n按键去抖\n模式/阈值设置"),
        ((560, 180, 930, 320), "#DCFCE7", "#16A34A", "DisplayTask\nOLED 显示\n亮度/温度/模式"),
        ((560, 390, 930, 530), "#FEF3C7", "#D97706", "ControlTask\n阈值判断\n舵机/LED/报警"),
        ((1060, 255, 1430, 415), "#EDE9FE", "#7C3AED", "WifiTask / DTU\nUSART1 JSON\n上云与命令解析"),
        ((1060, 500, 1430, 640), "#F8FAFC", "#94A3B8", "defaultTask\n系统空闲维护"),
    ]
    for xy, fill, outline, text in items:
        draw_round_box(draw, xy, fill, outline, text, body_font)
    arrow(draw, (430, 280), (560, 250), "#0284C7")
    arrow(draw, (430, 280), (560, 460), "#0284C7")
    arrow(draw, (930, 460), (1060, 335), "#D97706")
    arrow(draw, (430, 500), (560, 460), "#475569")
    draw_round_box(draw, (260, 700, 1240, 800), "#F8FAFC", "#CBD5E1",
                   "解耦价值：采集、显示、控制、通信互不阻塞；展示时可以独立验证真实传感器数据。",
                   small_font, radius=18)
    image.save(path)
    return path


def make_data_flow_diagram() -> Path:
    path = ASSET_DIR / "data_flow.png"
    image = Image.new("RGB", (1600, 900), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    title_font = pil_font(44)
    body_font = pil_font(26)
    small_font = pil_font(22)
    draw.text((70, 55), "真实数据采集与展示流程", font=title_font, fill="#0B2545")
    draw.text((70, 112), "页面不生成模拟数据，只解析来自板子的串口 JSON", font=small_font, fill="#64748B")

    steps = [
        ("传感器 AO", "光照/温度\n模拟电压"),
        ("ADC1", "PA0 / PA1\n12 位采样"),
        ("滤波换算", "16 点均值\n亮度% / 温度℃"),
        ("实时控制", "阈值判断\n舵机/LED/蜂鸣器"),
        ("JSON 输出", "[SENSOR]\nUSART2 每 5s"),
        ("演示台", "Web Serial\n趋势图/日志"),
    ]
    x = 70
    for idx, (title, desc) in enumerate(steps):
        draw_round_box(draw, (x, 260, x + 210, 440), "#FFFFFF", "#CBD5E1", f"{title}\n{desc}", body_font, radius=20)
        if idx < len(steps) - 1:
            arrow(draw, (x + 210, 350), (x + 270, 350), "#2563EB", width=5)
        x += 260

    sample = '{"light":68,"temp":27,"mode":0,"servo":2500,"sw1":1,"rssi":18}'
    draw_round_box(draw, (170, 610, 1430, 730), "#0F172A", "#0F172A", sample, pil_font(30), text_fill="#E2E8F0", radius=16)
    draw.text((210, 760), "比赛录制建议：遮挡光敏电阻、轻微加热热敏模块，观察 OLED、舵机、页面数值和趋势同步变化。",
              font=small_font, fill="#334155")
    image.save(path)
    return path


def create_project_plan(docx_path: Path, diagrams: dict[str, Path]) -> None:
    doc = Document()
    setup_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, PROJECT_TITLE, bold=True, color="0B2545", size=22)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(subtitle, PROJECT_SUBTITLE, color="4B5563", size=12)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(meta, "参赛方向：物联网终端 / 智能感知 / 嵌入式实时控制", color="64748B")
    add_run(meta, "\n团队/学校：兰州理工大学（可按实际报名信息补充）", color="64748B")
    add_run(meta, "\n版本：V1.0    日期：" + TODAY, color="64748B")
    doc.add_paragraph()
    add_note(
        doc,
        "定位说明",
        "本计划书根据当前工程代码、模块资料、银尔达平台联调记录与本地演示页面整理。已实现部分以代码和可录制现象为准；云平台物模型映射作为后续联调优化项单独说明，不夸大为已完全闭环。",
        "E0F2FE",
    )
    doc.add_page_break()

    doc.add_heading("执行概要", level=1)
    doc.add_paragraph(
        "“智境”面向室内学习、办公和实验空间，构建一套低成本、可扩展、可展示真实数据链路的环境智能感知与调控系统。系统以 STM32F103C8T6 为核心控制器，运行基于 CMSIS-RTOS v2 的 FreeRTOS 多任务程序，完成光照、温度等环境参数采集、滤波、显示、阈值判断、舵机/LED/蜂鸣器联动控制，并通过串口与 4G DTU、本地 Web Serial 仪表盘进行数据展示。"
    )
    doc.add_paragraph(
        "项目的核心价值不是简单读数，而是形成“感知—判断—执行—展示—上云”的完整物联网闭环雏形：传感器模块提供真实模拟量，FreeRTOS 任务拆分保证系统实时性和可维护性，OLED 和执行器提供现场反馈，本地可视化页面用于比赛录制时证明真实数据变化，银尔达 DTU/IoT 平台用于展示设备在线、SIM 激活和 4G 通信能力。"
    )
    add_note(
        doc,
        "当前可证明成果",
        "ADC 真实采样、OLED 显示、自动/手动控制、舵机/LED 联动、本地仪表盘 5 秒周期解析真实 JSON、银尔达平台设备/SIM 在线均已具备展示依据。平台物模型对传感器数据的最终解析仍需在银尔达侧继续调试数据流模板。",
        "F4F6F9",
    )

    doc.add_heading("项目背景", level=1)
    doc.add_heading("室内环境监测的现实需求", level=2)
    doc.add_paragraph(
        "教室、宿舍、实验室和办公空间长期存在光照不足、温度不适、通风与遮光调节滞后等问题。传统方案往往依赖人工感知和人工开关，缺少连续数据记录和自动调节能力；商业化智能家居系统成本较高，协议复杂，不利于学生竞赛项目快速搭建与扩展。"
    )
    doc.add_paragraph(
        "本项目选择光照与温度作为基础环境变量，一方面可以通过低成本模拟量传感器稳定采集，另一方面能够自然联动照明、遮光、通风、报警等控制场景。项目可作为室内环境调控、低功耗边缘节点、4G 远程监测和嵌入式实时系统教学的基础平台。"
    )
    doc.add_heading("痛点与机会", level=2)
    for item in [
        "痛点一：传感器读数与执行动作割裂，很多演示只停留在串口打印，缺乏直观可验证的调控效果。",
        "痛点二：上云配置复杂，设备在线不等于真实传感器数据已经进入平台物模型，容易在答辩中出现证据断层。",
        "痛点三：裸机循环难以表达物联网终端的任务并发、实时响应和扩展能力，缺少工程化说服力。",
        "机会：以 FreeRTOS 分层任务组织终端逻辑，用本地仪表盘和银尔达平台分别证明“真实数据”和“联网能力”，能形成稳定、可录制、可解释的比赛展示链路。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("产品与服务", level=1)
    doc.add_heading("系统定位", level=2)
    doc.add_paragraph(
        "智境是一套室内环境智能感知与调控原型系统，面向物联网设计大赛展示、嵌入式教学实践和低成本环境节点验证。它由嵌入式终端、执行模块、4G 通信模块、本地展示页面和云平台联调配置组成。"
    )
    doc.add_heading("核心功能模块", level=2)
    table = doc.add_table(rows=1, cols=4)
    headers = ["模块", "实现内容", "当前依据", "展示方式"]
    for i, header in enumerate(headers):
        add_run(table.cell(0, i).paragraphs[0], header, bold=True)
    rows = [
        ["环境采集", "光敏电阻 AO 接 PA0、热敏电阻 AO 接 PA1；ADC1 采样后 16 点滑动均值滤波。", "freertos.c 中 Task1、ADC_ReadChannel、SensorFilter。", "遮挡/加热传感器，观察 OLED 与仪表盘数值变化。"],
        ["现场显示", "OLED 显示亮度、温度、工作模式、阈值设置状态。", "DisplayTask 周期刷新 OLED。", "录制硬件屏幕和串口页面同步。"],
        ["自动调控", "光照低于阈值驱动报警 LED 与 SG90 舵机；温度高于阈值驱动状态 LED。", "ControlTask 读取全局传感器值并设置 TIM1 PWM。", "改变光照/温度，拍摄舵机/LED 动作。"],
        ["通信展示", "USART1 向 M100M-C2 DTU 输出 JSON；USART2 输出 [SENSOR] JSON 给本地仪表盘。", "Task4、DTU_SendData、SENSOR_DEBUG_INTERVAL=5000。", "本地页面 5 秒刷新真实数据；银尔达显示设备在线。"],
        ["人机交互", "KEY_MODE/KEY_SET 实现模式切换和阈值调整。", "KeyTask 外部中断与去抖逻辑。", "现场按键切换自动/手动/设置模式。"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            add_run(cells[i].paragraphs[0], value)
    style_table(table, [1.0, 2.35, 1.55, 1.9])

    doc.add_heading("核心技术方案", level=1)
    doc.add_heading("总体架构", level=2)
    doc.add_paragraph(
        "系统采用五层架构：感知层完成模拟量采集，控制层运行 STM32 + FreeRTOS，执行层负责 OLED、LED、蜂鸣器和舵机反馈，通信层通过 USART 与 M100M-C2 4G DTU 和本地调试口输出数据，展示层由本地 Web Serial 仪表盘和银尔达平台组成。"
    )
    doc.add_picture(str(diagrams["architecture"]), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("硬件设计与引脚规划", level=2)
    pin_table = doc.add_table(rows=1, cols=5)
    for i, header in enumerate(["资源", "STM32 引脚", "外设/模块", "方向", "说明"]):
        add_run(pin_table.cell(0, i).paragraphs[0], header, bold=True)
    pin_rows = [
        ["ADC1_IN0", "PA0-WKUP", "光敏传感器 AO", "输入", "采集环境光照原始模拟量"],
        ["ADC1_IN1", "PA1", "热敏传感器 AO", "输入", "采集环境温度原始模拟量"],
        ["TIM1_CH1", "PA8", "SG90 舵机 PWM", "输出", "执行遮光/开窗等调节动作"],
        ["USART1_TX", "PA9", "DTU RXD", "输出", "MCU 发送上行 JSON 到 4G DTU"],
        ["USART1_RX", "PA10", "DTU TXD", "输入", "接收平台/DTU 下行控制命令"],
        ["USART2_TX/RX", "PA2/PA3", "CH340/本地仪表盘", "双向", "调试日志与真实传感器 JSON 展示"],
        ["GPIO", "PB10/PB11", "DTU_RDY/DTU_RST", "输入/输出", "读取 DTU 连接状态、执行硬件复位"],
        ["GPIO", "PB8/PB9", "OLED SCL/SDA", "输出", "软件 I2C 驱动 OLED"],
        ["GPIO", "PA4/PB12", "KEY_MODE/KEY_SET", "输入", "模式切换与阈值设置"],
        ["GPIO", "PA6/PA7/PA11", "蜂鸣器/状态灯/报警灯", "输出", "现场报警与状态提示"],
    ]
    for row in pin_rows:
        cells = pin_table.add_row().cells
        for i, value in enumerate(row):
            add_run(cells[i].paragraphs[0], value)
    style_table(pin_table, [1.05, 1.05, 1.45, 0.7, 2.1])

    doc.add_heading("FreeRTOS 软件架构", level=2)
    doc.add_picture(str(diagrams["tasks"]), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "工程在 STM32CubeMX 中配置 CMSIS-RTOS v2，创建 SensorTask、DisplayTask、ControlTask、WifiTask、KeyTask 和 defaultTask 六个线程，并配置 xSensorQueue、xKeyQueue 两个消息队列。任务拆分使采集、显示、控制、通信互不阻塞，适合在答辩中解释实时系统设计思路。"
    )

    task_table = doc.add_table(rows=1, cols=4)
    for i, header in enumerate(["任务", "优先级/栈", "职责", "周期/触发"]):
        add_run(task_table.cell(0, i).paragraphs[0], header, bold=True)
    task_rows = [
        ["SensorTask", "AboveNormal / 256", "采集 PA0/PA1，滤波换算并更新全局真实传感器值。", "1s 采样；5s 输出 [SENSOR] JSON"],
        ["DisplayTask", "Normal / 256", "刷新 OLED，显示温度、亮度、模式和阈值设置界面。", "300ms"],
        ["ControlTask", "AboveNormal / 256", "根据阈值控制舵机、LED 和报警状态。", "200ms"],
        ["WifiTask", "BelowNormal / 512", "初始化 DTU，构造 JSON，处理下行命令。", "60s 上报"],
        ["KeyTask", "High / 128", "按键去抖、模式切换、阈值调节。", "按键事件触发"],
    ]
    for row in task_rows:
        cells = task_table.add_row().cells
        for i, value in enumerate(row):
            add_run(cells[i].paragraphs[0], value)
    style_table(task_table, [1.15, 1.0, 3.15, 1.25])

    doc.add_heading("数据采集、调控与通信流程", level=1)
    doc.add_picture(str(diagrams["data_flow"]), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("真实数据格式", level=2)
    doc.add_paragraph(
        "本地仪表盘推荐监听 USART2 调试口，解析固件每 5 秒输出一次的 [SENSOR] JSON。该数据来自 ADC 采样后的 g_light_value 和 g_temp_value，不由页面生成模拟值。"
    )
    add_note(
        doc,
        "本地展示 JSON 示例",
        '{"light":68,"temp":27,"mode":0,"servo":2500,"sw1":1,"in1":1,"rssi":18,"rdy_pin":1,"light_raw":1300,"temp_raw":1880,"fver":"V0.1","hver":"V1.0"}',
        "0F172A",
    )
    doc.add_heading("银尔达平台联调口径", level=2)
    doc.add_paragraph(
        "M100M-C2 4G DTU 已具备设备在线、SIM 激活、MQTT 通道配置和注册包成功记录。当前工程通过 USART1 发送扁平 JSON，理论上可由 DTU 透传至银尔达 IoT 平台；但从现有日志看，平台侧主要稳定收到 dreg/regbck 注册交互，传感器 dup 数据解析仍需继续优化数据流模板或任务脚本。因此比赛展示建议采用“双证据链”：银尔达平台展示设备在线与 4G 链路，本地仪表盘展示真实传感器数据。"
    )
    cloud_table = doc.add_table(rows=1, cols=3)
    for i, header in enumerate(["项目", "当前配置/现象", "比赛表达建议"]):
        add_run(cloud_table.cell(0, i).paragraphs[0], header, bold=True)
    cloud_rows = [
        ["服务器", "iot.yinerda.com:1883，MQTT 3.1.1，设备 ID/用户名使用平台设备 ID。", "说明系统已具备 4G 入网条件。"],
        ["注册包", '{"cmd":"dreg","did":"0","imei":"${IMEI}","iccid":"${ICCID}","pver":"test","fver":"V0.1"}', "展示日志中 dreg 成功解析。"],
        ["传感器上报", "MCU 已能在 USART1/USART2 输出 JSON；平台物模型映射仍需继续联调。", "不要把平台数据说成已稳定显示，强调可扩展和下一步。"],
        ["本地仪表盘", "Web Serial 只解析真实串口 JSON，5 秒节流刷新。", "作为视频里最稳定的数据证明页面。"],
    ]
    for row in cloud_rows:
        cells = cloud_table.add_row().cells
        for i, value in enumerate(row):
            add_run(cells[i].paragraphs[0], value)
    style_table(cloud_table, [1.0, 3.0, 2.2])

    doc.add_heading("创新点与竞争优势", level=1)
    for item in [
        "以 FreeRTOS 多任务而非裸机大循环组织终端逻辑，便于解释实时性、可维护性和后续扩展。",
        "将真实传感器数据、现场执行动作、OLED 显示和 Web 仪表盘绑定到同一数据源，演示可信度高。",
        "使用 4G DTU 而非仅依赖 Wi-Fi，适用于无局域网环境的远程物联网节点。",
        "在比赛展示上区分“真实数据证明”和“云平台在线证明”，降低平台物模型解析不稳定对演示的影响。",
        "为嘉立创 EDA 一体化载板预留清晰接口：STM32 最小系统、DTU、电源、传感器、OLED、舵机和调试口可集成为完整项目形态。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("实施计划与落地路径", level=1)
    plan_table = doc.add_table(rows=1, cols=5)
    for i, header in enumerate(["阶段", "目标", "关键任务", "验收指标", "状态"]):
        add_run(plan_table.cell(0, i).paragraphs[0], header, bold=True)
    plan_rows = [
        ["阶段 1", "终端采集与显示", "ADC、滤波、OLED、按键阈值、舵机/LED 联动。", "遮挡/加热后数值和执行器变化明显。", "已完成"],
        ["阶段 2", "本地真实数据演示", "USART2 输出 [SENSOR] JSON，Web Serial 仪表盘 5 秒刷新。", "页面只显示真实串口 JSON，最新上报不为空。", "已完成"],
        ["阶段 3", "4G 在线证明", "M100M-C2 供电、天线、SIM、MQTT 参数、注册包。", "银尔达 DTU/IoT 显示设备在线，日志注册成功。", "已完成"],
        ["阶段 4", "平台物模型完善", "调整数据流模板/任务脚本，使 dup 数据进入 temp/light 等功能点。", "IoT 设备数据页出现真实 temp/light。", "联调中"],
        ["阶段 5", "硬件集成", "嘉立创 EDA 设计底板，统一供电与接口丝印。", "模块直插，减少杜邦线，方便展示。", "计划中"],
    ]
    for row in plan_rows:
        cells = plan_table.add_row().cells
        for i, value in enumerate(row):
            add_run(cells[i].paragraphs[0], value)
    style_table(plan_table, [0.75, 1.1, 2.05, 1.75, 0.75])

    doc.add_heading("风险与对策", level=1)
    risk_table = doc.add_table(rows=1, cols=4)
    for i, header in enumerate(["风险", "表现", "影响", "对策"]):
        add_run(risk_table.cell(0, i).paragraphs[0], header, bold=True)
    risk_rows = [
        ["DTU 供电不足", "4G 发射瞬间掉电或反复离线", "云平台展示中断", "M100M-C2 使用 3.5~3.8V、足够电流电源，天线固定，录制前预热上线。"],
        ["平台物模型未解析", "日志只有 dreg/regbck，无 temp/light", "无法证明云端真实数据", "比赛展示采用本地仪表盘证明真实数据，同时说明平台映射为后续优化项。"],
        ["传感器噪声", "数值跳动大", "图表不稳定", "保留 16 点滑动均值，录制时操作幅度明显且间隔大于 5 秒。"],
        ["接线松动", "OLED/舵机/串口异常", "展示失败", "录制前固定杜邦线，后续通过嘉立创 EDA 载板集成。"],
        ["串口选择错误", "页面连接成功但没有 JSON", "本地仪表盘为空", "使用 PA2/PA3 的 USART2 调试口；若监听上行云数据则接 PA9。"],
    ]
    for row in risk_rows:
        cells = risk_table.add_row().cells
        for i, value in enumerate(row):
            add_run(cells[i].paragraphs[0], value)
    style_table(risk_table, [1.2, 1.45, 1.3, 2.4])

    doc.add_heading("经费与资源预算", level=1)
    budget_table = doc.add_table(rows=1, cols=4)
    for i, header in enumerate(["类别", "器件/资源", "估算金额", "说明"]):
        add_run(budget_table.cell(0, i).paragraphs[0], header, bold=True)
    budget_rows = [
        ["核心控制", "STM32F103C8T6 最小系统板、ST-LINK/CH340", "30~60 元", "已有套件可复用。"],
        ["感知与执行", "光敏、热敏、OLED、SG90、LED、蜂鸣器、按键", "40~80 元", "满足核心展示需求。"],
        ["通信", "银尔达 M100M-C2、SIM 卡、天线", "80~150 元", "用于 4G 入网展示。"],
        ["结构与 PCB", "嘉立创 EDA 载板、排针、端子、电源模块", "50~120 元", "后续集成减少接线风险。"],
        ["展示", "本地网页仪表盘、截图、PPT/计划书", "0 元", "软件自研。"],
    ]
    for row in budget_rows:
        cells = budget_table.add_row().cells
        for i, value in enumerate(row):
            add_run(cells[i].paragraphs[0], value)
    style_table(budget_table, [1.05, 2.25, 1.05, 2.45])

    doc.add_heading("预期成果与应用前景", level=1)
    doc.add_paragraph(
        "项目最终可以形成一套完整的物联网终端原型：硬件可插拔、程序可扩展、演示可验证、云端可继续接入。短期成果包括比赛视频、答辩 PPT、项目计划书和可运行硬件；中期可通过 PCB 载板、传感器扩展和平台物模型完善形成更接近产品的系统；长期可用于教室/实验室环境监测、智能窗帘/通风控制、校园低成本环境节点和嵌入式教学平台。"
    )

    doc.add_heading("比赛展示建议", level=1)
    for item in [
        "先展示整体硬件：STM32、传感器、OLED、舵机、DTU、CH340，说明接线关系。",
        "再展示本地仪表盘：连接 PA2/PA3 的 USART2，确认页面状态为“串口已连接”，等待 [SENSOR] JSON。",
        "录制遮挡光敏传感器：亮度下降，OLED/页面/舵机/报警灯同步变化。",
        "录制温度变化：热敏模块被手指或温源轻微加热，温度值和状态灯变化。",
        "最后切换银尔达平台：展示 SIM 在线、设备在线、注册日志成功，说明云端物模型映射是下一阶段完善重点。",
    ]:
        add_number(doc, item)

    doc.add_heading("项目总结", level=1)
    doc.add_paragraph(
        "智境项目以真实可运行的嵌入式终端为基础，通过 FreeRTOS 多任务架构把环境采集、现场显示、自动控制、4G 通信和可视化展示组织成一个完整系统。项目目前已经具备比赛录制所需的关键证据链：传感器真实数据、本地页面实时显示、执行器联动、4G 设备在线。后续通过银尔达平台数据模板完善和嘉立创 EDA 载板集成，可以从“竞赛原型”进一步提升为“可部署的室内环境智能节点”。"
    )

    screenshot_paths = [
        ROOT / "日志截图" / "设备日志.png",
        ROOT / "日志截图" / "设备数据.png",
    ]
    existing = [path for path in screenshot_paths if path.exists()]
    if existing:
        doc.add_heading("附录：联调截图", level=1)
        for path in existing:
            doc.add_paragraph(path.name)
            doc.add_picture(str(path), width=Inches(6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(docx_path)


VIDEO_SECTIONS = [
    {
        "time": "0:00-0:15",
        "scene": "开场标题",
        "action": "展示项目名称、整体硬件近景，镜头扫过 STM32、传感器、OLED、舵机和 DTU。",
        "script": "大家好，我们的项目是“智境——基于 FreeRTOS 的室内环境智能感知与调控系统”。它面向教室、宿舍和实验室等室内空间，实现环境数据采集、实时控制、可视化展示和 4G 物联网接入。",
        "evidence": "硬件完整性、项目名称清晰。",
    },
    {
        "time": "0:15-0:40",
        "scene": "问题与目标",
        "action": "切到 PPT 第 2 页或口播，说明室内环境监测痛点。",
        "script": "传统室内环境调节主要依赖人工判断，缺少连续数据和自动响应。我们希望用低成本嵌入式终端，把光照和温度这些环境变量转化为可显示、可控制、可上传的数据。",
        "evidence": "说明应用场景，不进入细枝末节。",
    },
    {
        "time": "0:40-1:10",
        "scene": "硬件架构",
        "action": "展示接线：PA0 光敏、PA1 热敏、PA8 舵机、PA9/PA10 DTU、PA2/PA3 本地串口。",
        "script": "硬件核心是 STM32F103C8T6。光敏和热敏模块通过 ADC 接入，OLED 显示现场数据，SG90 舵机和 LED 作为调控执行单元，M100M-C2 负责 4G 通信，本地 CH340 用于展示真实串口数据。",
        "evidence": "硬件连线与计划书一致。",
    },
    {
        "time": "1:10-1:45",
        "scene": "FreeRTOS 架构",
        "action": "展示 PPT 的任务架构页或代码片段，指向 SensorTask、DisplayTask、ControlTask、WifiTask、KeyTask。",
        "script": "软件上我们没有使用简单的裸机循环，而是用 FreeRTOS 拆成多个任务。SensorTask 采集传感器，DisplayTask 刷新 OLED，ControlTask 做阈值判断和执行控制，WifiTask 负责 DTU 通信，KeyTask 处理按键输入。这样每个功能互不阻塞，也方便后续扩展。",
        "evidence": "突出工程化和实时性。",
    },
    {
        "time": "1:45-2:35",
        "scene": "真实数据演示",
        "action": "打开本地仪表盘，连接串口，等待 5 秒周期数据；遮挡光敏传感器并观察数值和图表变化。",
        "script": "现在连接的是 STM32 的 USART2 调试口。页面不会生成模拟数据，只解析板子每 5 秒输出一次的真实 JSON。遮挡光敏传感器后，可以看到亮度数值下降，OLED 和页面同步变化，系统根据阈值驱动舵机和报警状态。",
        "evidence": "页面最新上报 JSON 非空，数值随物理操作变化。",
    },
    {
        "time": "2:35-3:05",
        "scene": "温度与执行动作",
        "action": "用手指或温源轻微接触热敏模块，等待 5 秒，拍摄温度变化和状态灯。",
        "script": "再看温度通道。热敏模块的模拟量经过 16 点滑动均值滤波后换算成温度值。当温度超过设定阈值时，状态灯会变化。这个过程体现了从感知到决策再到执行的闭环。",
        "evidence": "温度变化不要太快，等待页面刷新。",
    },
    {
        "time": "3:05-3:35",
        "scene": "4G/银尔达平台",
        "action": "切换到银尔达 DTU 或 IoT 平台，展示 SIM 在线、设备在线、设备日志 dreg/regbck 成功。",
        "script": "系统还接入了银尔达 M100M-C2 4G DTU。这里可以看到 SIM 卡和设备已经在线，注册包日志能够成功解析，说明终端具备蜂窝网络接入能力。当前平台物模型数据映射仍在继续完善，因此真实传感器数据展示以本地仪表盘作为主证据。",
        "evidence": "诚实说明云端联调状态，避免夸大。",
    },
    {
        "time": "3:35-4:10",
        "scene": "创新点与后续计划",
        "action": "展示 PPT 创新点、嘉立创 EDA 集成计划。",
        "script": "项目的创新点在于使用 FreeRTOS 组织多任务，把真实传感器、自动调控、本地可视化和 4G 联网能力组合成可扩展平台。下一步我们会用嘉立创 EDA 设计载板，减少杜邦线，提高稳定性，并完善银尔达平台物模型映射。",
        "evidence": "给出可落地后续路线。",
    },
    {
        "time": "4:10-4:25",
        "scene": "结尾",
        "action": "回到整体硬件画面和项目标题。",
        "script": "以上就是智境室内环境智能感知与调控系统的展示。它已经具备真实采集、现场控制、本地可视化和 4G 联网能力，后续可以继续扩展为校园和实验室环境节点。",
        "evidence": "收束明确。",
    },
]


def create_video_script_docx(docx_path: Path, md_path: Path) -> None:
    doc = Document()
    setup_doc(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, PROJECT_TITLE, bold=True, color="0B2545", size=20)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(subtitle, "讲解视频步骤与文案（建议 4 分 30 秒以内）", color="4B5563")
    add_note(
        doc,
        "录制原则",
        "更新速度不要太快，每一步操作后至少等待 5 秒，让页面显示的是板子真实输出的最新 JSON。不要使用模拟数据；若银尔达平台暂未显示 temp/light，就只展示在线和注册日志，把真实数据证明交给本地仪表盘。",
        "E0F2FE",
    )
    table = doc.add_table(rows=1, cols=5)
    for i, header in enumerate(["时间", "镜头/步骤", "操作", "讲解文案", "必须拍到的证据"]):
        add_run(table.cell(0, i).paragraphs[0], header, bold=True)
    for item in VIDEO_SECTIONS:
        cells = table.add_row().cells
        for i, key in enumerate(["time", "scene", "action", "script", "evidence"]):
            add_run(cells[i].paragraphs[0], item[key])
    style_table(table, [0.75, 1.0, 1.55, 2.15, 1.2])

    doc.add_heading("录制前检查清单", level=1)
    for item in [
        "M100M-C2 接好天线，供电稳定，SIM 在线。",
        "CH340 连接 STM32 的 USART2：PA2 接 CH340 RX，PA3 接 CH340 TX，共地，波特率 115200。",
        "本地仪表盘刷新间隔设为 5 秒，连接后等待 [SENSOR] JSON。",
        "光敏传感器遮挡和恢复都要慢一点，确保数值变化被采到。",
        "热敏模块不要用高温直接烤，手指或温和热源即可，避免损坏模块。",
        "银尔达平台截图只展示在线、注册、日志，不强行声称平台已稳定显示真实传感器数据。",
    ]:
        add_bullet(doc, item)
    doc.save(docx_path)

    lines = [
        f"# {PROJECT_TITLE}",
        "",
        "## 讲解视频步骤与文案",
        "",
        "录制原则：每一步操作后至少等待 5 秒，让页面显示板子的真实串口 JSON；不使用模拟数据。",
        "",
    ]
    for item in VIDEO_SECTIONS:
        lines.extend(
            [
                f"### {item['time']}｜{item['scene']}",
                f"- 操作：{item['action']}",
                f"- 文案：{item['script']}",
                f"- 证据：{item['evidence']}",
                "",
            ]
        )
    md_path.write_text("\n".join(lines), encoding="utf-8")


COMMON_MJS = r'''
export const C = {
  ink: "#0B172A",
  muted: "#64748B",
  blue: "#2563EB",
  cyan: "#0891B2",
  green: "#059669",
  amber: "#D97706",
  red: "#DC2626",
  purple: "#7C3AED",
  panel: "#FFFFFF",
  line: "#D9E2EC",
  bg: "#F7FAFC",
};

export function addBase(slide, ctx, title, kicker = "智境 · 室内环境智能感知与调控系统") {
  ctx.addShape(slide, { x: 0, y: 0, w: ctx.W, h: ctx.H, fill: C.bg, line: ctx.line("#00000000", 0) });
  ctx.addShape(slide, { x: 0, y: 0, w: ctx.W, h: 74, fill: "#FFFFFF", line: ctx.line(C.line, 1) });
  ctx.addText(slide, { text: kicker, x: 54, y: 20, w: 500, h: 28, fontSize: 18, bold: true, color: C.green, face: "Microsoft YaHei" });
  ctx.addText(slide, { text: title, x: 54, y: 86, w: 950, h: 58, fontSize: 34, bold: true, color: C.ink, face: "Microsoft YaHei" });
  ctx.addText(slide, { text: "物联网设计大赛", x: 1050, y: 24, w: 170, h: 25, fontSize: 16, color: C.muted, align: "right", face: "Microsoft YaHei" });
  ctx.addText(slide, { text: String(ctx.slideNumber).padStart(2, "0"), x: 1180, y: 675, w: 80, h: 28, fontSize: 15, color: C.muted, align: "right", face: "Microsoft YaHei" });
}

export function text(ctx, slide, value, x, y, w, h, size = 22, color = C.ink, bold = false, align = "left") {
  return ctx.addText(slide, { text: value, x, y, w, h, fontSize: size, color, bold, face: "Microsoft YaHei", insets: { left: 0, right: 0, top: 0, bottom: 0 }, align });
}

export function card(ctx, slide, x, y, w, h, title, body, accent = C.blue) {
  ctx.addShape(slide, { x, y, w, h, fill: C.panel, line: ctx.line(C.line, 1) });
  ctx.addShape(slide, { x, y, w: 8, h, fill: accent, line: ctx.line("#00000000", 0) });
  text(ctx, slide, title, x + 24, y + 22, w - 42, 30, 22, C.ink, true);
  text(ctx, slide, body, x + 24, y + 66, w - 42, h - 78, 18, C.muted);
}

export function chip(ctx, slide, label, x, y, w, fill = "#E0F2FE", color = C.blue) {
  ctx.addShape(slide, { x, y, w, h: 34, fill, line: ctx.line("#00000000", 0) });
  text(ctx, slide, label, x + 12, y + 7, w - 24, 22, 15, color, true, "center");
}

export function metric(ctx, slide, x, y, w, label, value, color = C.blue) {
  ctx.addShape(slide, { x, y, w, h: 112, fill: "#FFFFFF", line: ctx.line(C.line, 1) });
  text(ctx, slide, label, x + 18, y + 18, w - 36, 24, 16, C.muted);
  text(ctx, slide, value, x + 18, y + 48, w - 36, 42, 32, color, true);
}

export function flowBox(ctx, slide, x, y, w, h, title, body, color) {
  ctx.addShape(slide, { x, y, w, h, fill: "#FFFFFF", line: ctx.line(color, 2) });
  text(ctx, slide, title, x + 16, y + 16, w - 32, 26, 20, color, true, "center");
  text(ctx, slide, body, x + 16, y + 54, w - 32, h - 64, 17, C.muted, false, "center");
}

export function arrowText(ctx, slide, x, y) {
  text(ctx, slide, "→", x, y, 40, 40, 34, C.blue, true, "center");
}

export function codeBlock(ctx, slide, x, y, w, h, value) {
  ctx.addShape(slide, { x, y, w, h, fill: "#0F172A", line: ctx.line("#0F172A", 0) });
  ctx.addText(slide, { text: value, x: x + 18, y: y + 18, w: w - 36, h: h - 36, fontSize: 16, color: "#E2E8F0", face: "Consolas" });
}
'''


SLIDES: list[str] = []


def slide_module(number: int, body: str) -> str:
    return f"""import {{ addBase, text, card, chip, metric, flowBox, arrowText, codeBlock, C }} from './common.mjs';

export async function slide{number:02d}(presentation, ctx) {{
  const slide = presentation.slides.add();
{body}
  return slide;
}}
"""


def build_slide_sources() -> None:
    (SLIDES_DIR / "common.mjs").write_text(COMMON_MJS, encoding="utf-8")
    screenshots = {
        "log": (ROOT / "日志截图" / "设备日志.png").as_posix(),
        "data": (ROOT / "日志截图" / "设备数据.png").as_posix(),
        "arch": (ASSET_DIR / "system_architecture.png").as_posix(),
        "tasks": (ASSET_DIR / "freertos_task_architecture.png").as_posix(),
        "flow": (ASSET_DIR / "data_flow.png").as_posix(),
    }
    slide_bodies = [
        r'''
  addBase(slide, ctx, "智境 — 基于 FreeRTOS 的室内环境智能感知与调控系统", "SMART ENVIRONMENT MONITOR");
  text(ctx, slide, "面向室内空间的真实传感器采集、实时控制与 4G 物联网接入原型", 54, 168, 820, 44, 24, C.muted);
  metric(ctx, slide, 54, 260, 230, "核心控制器", "STM32F103", C.blue);
  metric(ctx, slide, 310, 260, 230, "实时系统", "FreeRTOS", C.green);
  metric(ctx, slide, 566, 260, 230, "通信能力", "4G DTU", C.purple);
  metric(ctx, slide, 822, 260, 230, "真实展示", "Web Serial", C.cyan);
  card(ctx, slide, 54, 466, 550, 132, "项目一句话", "用低成本 STM32 终端持续采集室内光照和温度，按阈值自动驱动舵机/LED/报警，并把真实数据送到本地展示与云端链路。", C.green);
  card(ctx, slide, 630, 466, 550, 132, "答辩口径", "真实数据以本地仪表盘和 OLED/执行器联动为主证据；银尔达平台展示设备在线、SIM 在线和 4G 链路能力。", C.blue);
''',
        r'''
  addBase(slide, ctx, "为什么做：室内环境调节需要可验证的闭环");
  card(ctx, slide, 74, 170, 340, 240, "痛点 1：靠人工感知", "光照不足、温度不适往往靠人主观判断，缺少连续采样和自动响应。", C.red);
  card(ctx, slide, 460, 170, 340, 240, "痛点 2：演示证据断层", "设备在线不等于真实传感器数据已进入平台；比赛中需要可录制、可复现的证据链。", C.amber);
  card(ctx, slide, 846, 170, 340, 240, "痛点 3：裸机扩展受限", "单循环逻辑难以说明实时响应、任务并发和后续功能扩展。", C.purple);
  text(ctx, slide, "目标：构建“感知 → 判断 → 执行 → 展示 → 上云”的室内环境物联网终端原型", 110, 500, 1030, 42, 28, C.ink, true, "center");
  chip(ctx, slide, "真实 ADC 数据", 196, 568, 160, "#DCFCE7", C.green);
  chip(ctx, slide, "FreeRTOS 多任务", 402, 568, 190, "#E0F2FE", C.blue);
  chip(ctx, slide, "执行器联动", 640, 568, 150, "#FEF3C7", C.amber);
  chip(ctx, slide, "4G 在线能力", 840, 568, 160, "#EDE9FE", C.purple);
''',
        f'''
  addBase(slide, ctx, "系统总体架构：把模块组织成项目");
  await ctx.addImage(slide, {{ path: "{screenshots['arch']}", x: 70, y: 150, w: 1140, h: 530, fit: "contain", alt: "系统总体架构图" }});
''',
        r'''
  addBase(slide, ctx, "硬件方案：低成本模块 + 清晰引脚映射");
  flowBox(ctx, slide, 70, 168, 170, 120, "PA0", "光敏 AO\nADC1_IN0", C.blue);
  flowBox(ctx, slide, 270, 168, 170, 120, "PA1", "热敏 AO\nADC1_IN1", C.blue);
  flowBox(ctx, slide, 470, 168, 170, 120, "PA8", "TIM1_CH1\nSG90 PWM", C.amber);
  flowBox(ctx, slide, 670, 168, 210, 120, "PA9 / PA10", "USART1\n4G DTU RX/TX", C.purple);
  flowBox(ctx, slide, 910, 168, 210, 120, "PA2 / PA3", "USART2\n本地仪表盘", C.cyan);
  card(ctx, slide, 70, 340, 360, 188, "现场交互", "OLED 使用 PB8/PB9 软件 I2C；KEY_MODE/KEY_SET 用于模式切换和阈值设置；LED 与蜂鸣器给出状态反馈。", C.green);
  card(ctx, slide, 470, 340, 360, 188, "4G 通信", "M100M-C2 供电建议 3.5~3.8V，RXD 接 PA9，TXD 接 PA10，RDY/RST 接 PB10/PB11。", C.purple);
  card(ctx, slide, 870, 340, 360, 188, "PCB 集成方向", "嘉立创 EDA 载板预留 STM32、DTU、传感器、OLED、舵机、电源和调试接口，减少杜邦线风险。", C.blue);
  text(ctx, slide, "关键原则：TX/RX 交叉、所有模块共地、DTU 供电留足电流余量。", 100, 590, 1040, 34, 24, C.ink, true, "center");
''',
        f'''
  addBase(slide, ctx, "FreeRTOS 软件架构：采集、显示、控制、通信解耦");
  await ctx.addImage(slide, {{ path: "{screenshots['tasks']}", x: 72, y: 150, w: 1136, h: 520, fit: "contain", alt: "FreeRTOS任务架构图" }});
''',
        f'''
  addBase(slide, ctx, "真实数据链路：页面只接收板子 JSON");
  await ctx.addImage(slide, {{ path: "{screenshots['flow']}", x: 64, y: 148, w: 1160, h: 475, fit: "contain", alt: "真实数据流程图" }});
  chip(ctx, slide, "刷新间隔：5 秒", 980, 622, 160, "#DCFCE7", C.green);
''',
        r'''
  addBase(slide, ctx, "上报格式：扁平 JSON，便于平台与本地解析");
  codeBlock(ctx, slide, 74, 172, 760, 210, `{"light":68,"temp":27,"mode":0,"servo":2500,
"sw1":1,"in1":1,"rssi":18,"iccid":"8986...",
"imei":"8622...","fver":"V0.1","hver":"V1.0"}`);
  card(ctx, slide, 880, 172, 300, 210, "字段含义", "light/temp 为真实传感器值；mode 表示自动/手动；servo 为 PWM；sw1 为报警状态；rssi 为信号强度。", C.blue);
  card(ctx, slide, 74, 430, 360, 170, "本地仪表盘", "监听 USART2 的 [SENSOR] JSON，每 5 秒刷新，适合比赛录制证明真实数据。", C.green);
  card(ctx, slide, 474, 430, 360, 170, "银尔达平台", "设备/SIM 在线、注册包成功；传感器物模型映射仍需继续联调数据流模板。", C.purple);
  card(ctx, slide, 874, 430, 300, 170, "答辩重点", "不把平台数据缺口夸大，明确展示双证据链。", C.amber);
''',
        f'''
  addBase(slide, ctx, "联调证据：设备在线 + 注册日志成功");
  await ctx.addImage(slide, {{ path: "{screenshots['log']}", x: 62, y: 155, w: 555, h: 392, fit: "contain", alt: "银尔达设备日志截图" }});
  await ctx.addImage(slide, {{ path: "{screenshots['data']}", x: 650, y: 155, w: 555, h: 392, fit: "contain", alt: "银尔达设备数据截图" }});
  text(ctx, slide, "云平台证明联网能力；真实传感器变化通过本地仪表盘和 OLED/执行器同步证明。", 110, 585, 1060, 40, 25, C.ink, true, "center");
''',
        r'''
  addBase(slide, ctx, "比赛演示：四步证明系统可用");
  flowBox(ctx, slide, 80, 180, 230, 150, "1. 连接串口", "Chrome/Edge 打开本地仪表盘\n波特率 115200", C.cyan);
  arrowText(ctx, slide, 328, 235);
  flowBox(ctx, slide, 380, 180, 230, 150, "2. 等待 JSON", "页面 5 秒刷新\n最新上报不为空", C.green);
  arrowText(ctx, slide, 628, 235);
  flowBox(ctx, slide, 680, 180, 230, 150, "3. 改变环境", "遮挡光敏\n轻微加热热敏", C.amber);
  arrowText(ctx, slide, 928, 235);
  flowBox(ctx, slide, 980, 180, 230, 150, "4. 观察联动", "OLED + 图表 + 舵机/LED\n同步变化", C.blue);
  card(ctx, slide, 120, 420, 500, 160, "录制注意", "不要快速连续操作；每次改变环境后等待至少 5 秒，让页面显示真实串口数据。", C.red);
  card(ctx, slide, 680, 420, 500, 160, "平台说明", "若银尔达未显示 temp/light，只展示设备在线和日志成功，不强行解释为云端数据闭环。", C.purple);
''',
        r'''
  addBase(slide, ctx, "创新点：从模块实验走向项目系统");
  card(ctx, slide, 76, 170, 350, 190, "工程化实时系统", "FreeRTOS 将采集、显示、控制、通信拆分为独立任务，展示的不只是传感器读数。", C.green);
  card(ctx, slide, 466, 170, 350, 190, "真实数据可信展示", "本地仪表盘只解析板子 JSON，并以 5 秒节流避免页面更新过快或误读模拟值。", C.blue);
  card(ctx, slide, 856, 170, 350, 190, "4G 物联网能力", "M100M-C2 已完成 SIM 在线和设备在线，可扩展到远程监测场景。", C.purple);
  card(ctx, slide, 270, 430, 350, 160, "可集成硬件", "后续用嘉立创 EDA 绘制模块载板，统一供电、串口和传感器接口。", C.amber);
  card(ctx, slide, 660, 430, 350, 160, "可扩展平台", "可增加 CO₂、湿度、人体存在、继电器、风扇等模块。", C.cyan);
''',
        r'''
  addBase(slide, ctx, "实施计划：当前可演示，后续可产品化");
  metric(ctx, slide, 78, 168, 220, "阶段 1", "终端闭环", C.green);
  metric(ctx, slide, 322, 168, 220, "阶段 2", "本地展示", C.blue);
  metric(ctx, slide, 566, 168, 220, "阶段 3", "4G 在线", C.purple);
  metric(ctx, slide, 810, 168, 220, "阶段 4", "平台映射", C.amber);
  metric(ctx, slide, 1054, 168, 220, "阶段 5", "PCB 集成", C.cyan);
  card(ctx, slide, 80, 372, 360, 170, "已完成", "ADC 采样、OLED 显示、舵机/LED 控制、本地仪表盘真实 JSON、银尔达在线。", C.green);
  card(ctx, slide, 480, 372, 360, 170, "联调中", "银尔达 IoT 物模型 temp/light 数据解析，继续优化数据流模板和任务脚本。", C.amber);
  card(ctx, slide, 880, 372, 360, 170, "下一步", "嘉立创 EDA 载板、供电稳定性、结构固定、演示外观整理。", C.blue);
''',
        r'''
  addBase(slide, ctx, "讲解视频脚本：控制在 4 分 30 秒以内");
  card(ctx, slide, 72, 158, 360, 150, "0:00-0:40 开场与痛点", "项目名称、整体硬件、室内环境监测需求。", C.green);
  card(ctx, slide, 470, 158, 360, 150, "0:40-1:45 架构说明", "硬件引脚、FreeRTOS 任务分工、真实数据来源。", C.blue);
  card(ctx, slide, 868, 158, 360, 150, "1:45-3:05 真实演示", "连接串口、等待 JSON、遮挡光敏、温度变化、执行器联动。", C.amber);
  card(ctx, slide, 270, 390, 360, 150, "3:05-3:35 云平台", "展示银尔达在线、SIM 在线、注册日志成功。", C.purple);
  card(ctx, slide, 668, 390, 360, 150, "3:35-4:25 总结", "创新点、嘉立创 EDA 集成计划、应用前景。", C.cyan);
  text(ctx, slide, "关键口径：真实数据看本地仪表盘；联网能力看银尔达平台。", 160, 598, 960, 38, 26, C.ink, true, "center");
''',
        r'''
  addBase(slide, ctx, "结论：智境已具备可录制、可解释、可扩展的比赛展示能力");
  text(ctx, slide, "感知", 120, 180, 170, 46, 34, C.blue, true, "center");
  text(ctx, slide, "控制", 330, 180, 170, 46, 34, C.green, true, "center");
  text(ctx, slide, "展示", 540, 180, 170, 46, 34, C.amber, true, "center");
  text(ctx, slide, "联网", 750, 180, 170, 46, 34, C.purple, true, "center");
  text(ctx, slide, "集成", 960, 180, 170, 46, 34, C.cyan, true, "center");
  card(ctx, slide, 160, 300, 960, 170, "答辩结束语", "项目已经完成真实传感器采集、本地可视化、现场执行控制和 4G 在线能力验证。后续通过平台物模型完善和 PCB 载板集成，可以继续向可部署的室内环境智能节点演进。", C.green);
  text(ctx, slide, "谢谢观看 / Q&A", 430, 545, 420, 56, 38, C.ink, true, "center");
''',
    ]
    for idx, body in enumerate(slide_bodies, 1):
        (SLIDES_DIR / f"slide-{idx:02d}.mjs").write_text(slide_module(idx, body), encoding="utf-8")

    profile_plan = """task mode: create
primary deck-profile: engineering-platform
secondary gates: product-platform, competition demo narrative
required proof objects:
- architecture map
- FreeRTOS task map
- data-flow and JSON proof
- online-platform screenshots
- demo procedure and roadmap
known missing inputs:
- team member names and competition registration fields should be filled by the user if required.
"""
    (PRESENTATION_WORKSPACE / "profile-plan.txt").write_text(profile_plan, encoding="utf-8")
    (PRESENTATION_WORKSPACE / "source-notes.txt").write_text(
        "Sources: local STM32 project files, extracted template structure, module notes, Yinerda screenshots/log screenshots, local dashboard source.\n",
        encoding="utf-8",
    )


def write_material_summary(paths: dict[str, Path]) -> None:
    summary = {
        "project_title": PROJECT_TITLE,
        "outputs": {key: str(value) for key, value in paths.items()},
        "truth_statement": "真实传感器数据以 USART2 [SENSOR] JSON、本地仪表盘、OLED 和执行器联动为主证据；银尔达平台用于展示设备在线和 4G 链路。",
        "recommended_demo_interval_seconds": 5,
    }
    (OUT_DIR / "材料生成摘要.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    ensure_dirs()
    diagrams = {
        "architecture": make_architecture_diagram(),
        "tasks": make_task_diagram(),
        "data_flow": make_data_flow_diagram(),
    }
    project_plan = OUT_DIR / "智境-项目计划书.docx"
    video_docx = OUT_DIR / "智境-讲解视频步骤与文案.docx"
    video_md = OUT_DIR / "智境-讲解视频步骤与文案.md"
    create_project_plan(project_plan, diagrams)
    create_video_script_docx(video_docx, video_md)
    build_slide_sources()
    write_material_summary(
        {
            "project_plan_docx": project_plan,
            "video_script_docx": video_docx,
            "video_script_markdown": video_md,
            "presentation_pptx": OUT_DIR / "智境-物联网设计大赛答辩PPT.pptx",
            "ppt_workspace": PRESENTATION_WORKSPACE,
            "slides_dir": SLIDES_DIR,
        }
    )
    print(json.dumps({
        "project_plan": str(project_plan),
        "video_docx": str(video_docx),
        "video_md": str(video_md),
        "slides_dir": str(SLIDES_DIR),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
